import openpyxl
wb = openpyxl.load_workbook("D:/Skill Rubri/agenteval_test01/data/eval_results_cn.xlsx")
ws = wb["论文评测"]
headers = [cell.value for cell in ws[1]]
print("Excel 列名:", headers)
print()
row2 = [cell.value for cell in ws[2]]
for h, v in zip(headers, row2):
    print(f"{h}: {str(v)[:80]}")
